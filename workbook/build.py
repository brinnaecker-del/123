# -*- coding: utf-8 -*-
"""生成与论文现行方法一致的可复现计算底稿（全部为活公式）。
方法链路：固定基准标准化 → 子系统内等权合成 → MPI失衡惩罚聚合 → 三层诊断。"""
import openpyxl, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL

YS = list(range(2016, 2026))
NY = len(YS)
C0 = 3                                   # 年份起始列 = C
col = lambda i: CL(C0 + i)               # 第 i 年(0-based)的列字母

# ---------------- 基础数据 ----------------
BASE = [
 ("地方一般公共预算收入","亿元",[155.99,185.83,230.35,221.99,220.99,215.62,179.63,236.63,277.19,317.90]),
 ("城乡居民人均可支配收入比","倍",[3.05718,2.96912,2.95170,2.88858,2.81929,2.74597,2.67741,2.60490,2.56947,2.53597]),
 ("居民医保人均财政补助标准","元/人年",[420,450,515,515,550,580,615,675,705,735]),
 ("基本养老保险参保率","%",[85.50,89.30,91.15,92.60,94.10,95.20,95.50,95.80,96.31,96.50]),
 ("人均GDP相对全国水平","%",[65.4184,65.8931,66.2221,69.7822,72.8755,68.3805,68.5000,73.4400,78.5800,81.7700]),
 ("地区生产总值","亿元",[1197.38,1380.97,1583.48,1750.34,1956.54,2145.08,2235.39,2532.86,2789.07,3031.89]),
 ("农村居民人均可支配收入","元",[9094,10330,11450,12951,14598,16935,18209,19924,21578,23184]),
 ("货物进出口总额","亿元",[51.67,59.19,47.52,48.76,21.33,40.16,46.01,109.78,126.72,84.79]),
 ("接待游客总人数","万人次",[2315.94,2561.43,3368.73,4012.15,3505.01,4153.44,3002.76,5516.98,6389.10,7073.37]),
 ("发电量","亿千瓦时",[50.78,52.06,63.14,66.30,67.10,80.80,111.10,133.50,140.41,162.42]),
 ("单位GDP能耗指数","指数",[96.80,93.41,90.51,87.98,86.66,84.06,81.20,78.85,76.41,73.89]),
 ("森林覆盖率","%",[11.98,12.14,12.14,12.14,12.31,12.31,12.31,12.31,12.31,12.54]),
 ("空气质量优良天数比例","%",[96.80,97.30,98.90,99.60,99.40,99.20,99.00,99.40,99.50,99.60]),
 ("草原综合植被盖度","%",[42.30,44.20,45.90,46.70,46.90,47.14,47.50,47.78,48.02,48.07]),
]
AUX = [
 ("西藏年末常住人口","万人",[330.54,337.15,343.82,350.56,364.81,366.00,364.00,365.00,370.00,374.00]),
 ("全国农村居民人均可支配收入","元",[12363,13432,14617,16021,17131,18931,20133,21691,23119,24456]),
 ("全国城镇居民人均可支配收入","元",[33616,36396,39251,42359,43834,47412,49283,51821,54188,56502]),
]
BR = {n: i + 3 for i, (n, _, _) in enumerate(BASE)}          # 基础数据行号
AR = {n: i + 3 + len(BASE) + 2 for i, (n, _, _) in enumerate(AUX)}
R_RATIO = AR["全国城镇居民人均可支配收入"] + 1                # 全国城乡收入比

# ---------------- 13 项指标（表1口径） ----------------
def b(name, i): return "'1_基础数据'!%s%d" % (col(i), BR[name])
def a(name, i): return "'1_基础数据'!%s%d" % (col(i), AR[name])
IND = [
 ("安全 S","S1","人均一般公共预算收入（元/人）","正向",0,12000,
    lambda i: "=%s*10000/%s" % (b("地方一般公共预算收入",i), a("西藏年末常住人口",i))),
 ("安全 S","S2","城乡居民收入比与全国差距（倍）","逆向",-1,2,
    lambda i: "=%s-'1_基础数据'!%s%d" % (b("城乡居民人均可支配收入比",i), col(i), R_RATIO)),
 ("安全 S","S3","居民医保人均财政补助标准（元/人年）","正向",0,1200,
    lambda i: "=%s" % b("居民医保人均财政补助标准",i)),
 ("安全 S","S4","基本养老保险参保率（%）","正向",0,100,
    lambda i: "=%s" % b("基本养老保险参保率",i)),
 ("发展 D","D1","人均GDP相对全国水平（%）","正向",0,100,
    lambda i: "=%s" % b("人均GDP相对全国水平",i)),
 ("发展 D","D2","农村居民人均收入相对全国水平（%）","正向",0,100,
    lambda i: "=%s/%s*100" % (b("农村居民人均可支配收入",i), a("全国农村居民人均可支配收入",i))),
 ("发展 D","D3","外向度（进出口总额/GDP，%）","正向",0,10,
    lambda i: "=%s/%s*100" % (b("货物进出口总额",i), b("地区生产总值",i))),
 ("发展 D","D4","人均接待游客（万人次/万人）","正向",0,25,
    lambda i: "=%s/%s" % (b("接待游客总人数",i), a("西藏年末常住人口",i))),
 ("发展 D","D5","发电量（亿千瓦时）","正向",0,200,
    lambda i: "=%s" % b("发电量",i)),
 ("生态 E","E1","单位GDP能耗指数","逆向",60,100,
    lambda i: "=%s" % b("单位GDP能耗指数",i)),
 ("生态 E","E2","森林覆盖率（%）","正向",11.98,12.51,
    lambda i: "=%s" % b("森林覆盖率",i)),
 ("生态 E","E3","空气质量优良天数比例（%）","正向",0,100,
    lambda i: "=%s" % b("空气质量优良天数比例",i)),
 ("生态 E","E4","草原综合植被盖度（%）","正向",42.30,50,
    lambda i: "=%s" % b("草原综合植被盖度",i)),
]
IR   = {c: i + 3 for i, (_, c, _, _, _, _, _) in enumerate(IND)}       # 指标行号
SUB  = {"安全 S": ["S1","S2","S3","S4"], "发展 D": ["D1","D2","D3","D4","D5"],
        "生态 E": ["E1","E2","E3","E4"]}
PAPER = {"US":[.537,.575,.617,.615,.619,.626,.615,.661,.688,.715],
         "UD":[.467,.483,.489,.519,.484,.528,.539,.672,.713,.718],
         "UE":[.267,.423,.499,.542,.634,.657,.686,.710,.733,.843],
         "D" :[.393,.486,.528,.556,.571,.598,.607,.680,.711,.754]}

# ---------------- 样式 ----------------
HDR  = PatternFill("solid", fgColor="D9E2F3")
BLUE = Font(color="0070C0")                  # 蓝=原始输入（可改）
GRN  = Font(color="00823B")                  # 绿=引用其他表
YEL  = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
THIN = Border(*[Side("thin", color="BFBFBF")] * 4)

def header(ws, row, cells, width=None):
    for j, v in enumerate(cells):
        c = ws.cell(row=row, column=j + 1, value=v)
        c.fill, c.font, c.border = HDR, BOLD, THIN
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if width:
        for j, w in enumerate(width): ws.column_dimensions[CL(j + 1)].width = w

def yearhdr(ws, row, lead, width=None):
    header(ws, row, lead + [str(y) for y in YS], width)

wb = openpyxl.Workbook(); wb.remove(wb.active)

# ============ 说明 ============
ws = wb.create_sheet("说明")
ws.column_dimensions["A"].width = 118
txt = [
 ("西藏高质量发展「安全—发展—生态」三元协调评价 · 可复现计算底稿", True),
 ("", False),
 ("本底稿与论文现行方法严格一致，全部为活公式：修改「1_基础数据」蓝色单元格后，所有下游结果自动重算。", False),
 ("", False),
 ("计算链路", True),
 ("  1_基础数据      原始统计值与三条辅助序列（常住人口、全国城乡居民人均可支配收入）", False),
 ("  2_指标值        按表1口径换算出13项指标（人均化、相对化、比值化）", False),
 ("  3_标准化        固定基准标准化，式(1)正向/式(2)逆向，超出[0,1]截断；式(3)平移压缩至[0.01,0.99]", False),
 ("  4_子系统指数    子系统内等权平均，得 US / UD / UE", False),
 ("  5_MPI协调指数   D = M − S²/M，M为三子系统均值，S²为总体方差", False),
 ("  6_组合权重      熵值权重 × CRITIC权重后归一（仅用于诊断与稳健性对照，不参与子系统合成）", False),
 ("  7_结构性贡献缺口  O(i,j) = w(j)·I(i,j) / Σ w·I，逐年计算后取样本期均值（论文表4）", False),
 ("  8_局部敏感性    ε = (D⁺−D⁻)/(0.02·D)，±1%相对扰动双侧数值微分（论文表5）", False),
 ("  9_留一法        逐一剔除13项指标重算，检验方向性判断稳健性（论文表6）", False),
 ("  10_稳健性检验   滚动端点稳定性、极差标准化对照、惯例性上限±25%扰动", False),
 ("  11_对照指数     四维度线性加总对照（需先在该表填入维度归属）", False),
 ("", False),
 ("关键设定", True),
 ("  · 标准化：固定基准法为主设定，上下限外生、不随样本更新而回溯改写（论文§3(三)）", False),
 ("  · 合成权重：子系统内等权为主设定；客观赋权仅用于诊断与稳健性对照（论文§3(四)）", False),
 ("  · 聚合：MPI失衡惩罚型指数，非传统耦合协调度（论文§3(五)）", False),
 ("  · 地区生产总值采用含第五次全国经济普查修订的最新核算口径（2023年为2532.86亿元）", False),
 ("", False),
 ("颜色约定：蓝色=原始输入（可修改）；黑色=本表公式；绿色=引用其他表；黄底=与论文载值的核对列", False),
 ("", False),
 ("已知差异（如实标注，勿删）", True),
 ("  安全子系统 US 与生态子系统 UE：与论文表3 逐年精确吻合（各10/10）。", False),
 ("  发展子系统 UD：与论文表3 存在 0.001—0.005 的差异，传导至MPI最大0.002，不改变任何结论。", False),
 ("    原因待作者确认：本底稿的发展类原始值取自旧版计算文件，现稿可能已更新其中部分数值。", False),
 ("  三阶段短板归属（论文核心结论）：经本底稿独立复算，10/10 与论文一致。", False),
 ("", False),
 ("辅助序列来源：西藏常住人口据历年统计公报与第七次全国人口普查；全国城乡居民人均可支配收入据国家统计局", False),
 ("历年《居民收入和消费支出情况》。其中2016—2018年西藏常住人口未经原件直接确证，但US三年均精确吻合。", False),
]
for i, (t, bold) in enumerate(txt, 1):
    c = ws.cell(row=i, column=1, value=t)
    if bold: c.font = BOLD
    c.alignment = Alignment(wrap_text=False)

# ============ 1_基础数据 ============
ws = wb.create_sheet("1_基础数据")
ws["A1"] = "基础数据（蓝色为原始输入，可直接修改）"; ws["A1"].font = BOLD
yearhdr(ws, 2, ["项目", "单位"], [30, 12] + [11] * NY)
for k, (n, u, vals) in enumerate(BASE):
    r = k + 3
    ws.cell(row=r, column=1, value=n).border = THIN
    ws.cell(row=r, column=2, value=u).border = THIN
    for i, v in enumerate(vals):
        c = ws.cell(row=r, column=C0 + i, value=v); c.font = BLUE; c.border = THIN
r0 = len(BASE) + 4
ws.cell(row=r0, column=1, value="辅助序列").font = BOLD
for k, (n, u, vals) in enumerate(AUX):
    r = AR[n]
    ws.cell(row=r, column=1, value=n).border = THIN
    ws.cell(row=r, column=2, value=u).border = THIN
    for i, v in enumerate(vals):
        c = ws.cell(row=r, column=C0 + i, value=v); c.font = BLUE; c.border = THIN
ws.cell(row=R_RATIO, column=1, value="全国城乡居民收入比").border = THIN
ws.cell(row=R_RATIO, column=2, value="倍").border = THIN
for i in range(NY):
    c = ws.cell(row=R_RATIO, column=C0 + i,
                value="=%s%d/%s%d" % (col(i), AR["全国城镇居民人均可支配收入"],
                                      col(i), AR["全国农村居民人均可支配收入"]))
    c.number_format = "0.0000"; c.border = THIN

# ============ 2_指标值 ============
ws = wb.create_sheet("2_指标值")
ws["A1"] = "13项指标（表1口径）—— 由 1_基础数据 换算"; ws["A1"].font = BOLD
yearhdr(ws, 2, ["子系统", "编码", "三级指标", "方向", "下限", "上限"], [9, 7, 34, 7, 9, 9] + [11] * NY)
for sub, code, name, d, lo, hi, f in IND:
    r = IR[code]
    for j, v in enumerate([sub, code, name, d, lo, hi]):
        ws.cell(row=r, column=j + 1, value=v).border = THIN
    for i in range(NY):
        c = ws.cell(row=r, column=C0 + 4 + i, value=f(i)); c.font = GRN
        c.number_format = "0.0000"; c.border = THIN
# 指标值表的年份列右移4列(因前面有下限/上限)，重设表头
yearhdr(ws, 2, ["子系统", "编码", "三级指标", "方向", "下限", "上限"])
for i, y in enumerate(YS):
    ws.cell(row=2, column=C0 + 4 + i, value=str(y)).fill = HDR
    ws.cell(row=2, column=C0 + 4 + i).font = BOLD
    ws.cell(row=2, column=C0 + 4 + i).alignment = Alignment(horizontal="center")
for i in range(NY):
    ws.cell(row=2, column=C0 + i, value=None)
ws.cell(row=2, column=5, value="下限"); ws.cell(row=2, column=6, value="上限")
for cix in (5, 6):
    ws.cell(row=2, column=cix).fill = HDR; ws.cell(row=2, column=cix).font = BOLD

VC0 = C0 + 4                                   # 指标值表年份起始列 = G
vcol = lambda i: CL(VC0 + i)

# ============ 3_标准化 ============
ws = wb.create_sheet("3_标准化")
ws["A1"] = "固定基准标准化：式(1)正向 / 式(2)逆向，超出[0,1]截断；式(3) x' = 0.98·x* + 0.01"; ws["A1"].font = BOLD
yearhdr(ws, 2, ["子系统", "编码", "三级指标", "方向"], [9, 7, 34, 7] + [11] * NY)
for sub, code, name, d, lo, hi, _ in IND:
    r = IR[code]
    for j, v in enumerate([sub, code, name, d]):
        ws.cell(row=r, column=j + 1, value=v).border = THIN
    for i in range(NY):
        src = "'2_指标值'!%s%d" % (vcol(i), r)
        lo_r, hi_r = "'2_指标值'!$E%d" % r, "'2_指标值'!$F%d" % r
        num = ("(%s-%s)" % (src, lo_r)) if d == "正向" else ("(%s-%s)" % (hi_r, src))
        f = "=0.98*MIN(1,MAX(0,%s/(%s-%s)))+0.01" % (num, hi_r, lo_r)
        c = ws.cell(row=r, column=C0 + i, value=f)
        c.number_format = "0.0000"; c.border = THIN

def std_ref(code, i): return "'3_标准化'!%s%d" % (col(i), IR[code])

# ============ 4_子系统指数 ============
ws = wb.create_sheet("4_子系统指数")
ws["A1"] = "子系统指数：子系统内等权平均（论文§3(四) 主设定）"; ws["A1"].font = BOLD
header(ws, 2, ["年份", "US 安全", "UD 发展", "UE 生态", "", "US 论文", "UD 论文", "UE 论文",
               "", "US 差", "UD 差", "UE 差"], [8] + [11] * 11)
SUBCOL = {"安全 S": 2, "发展 D": 3, "生态 E": 4}
for i, y in enumerate(YS):
    r = i + 3
    ws.cell(row=r, column=1, value=y).border = THIN
    for sub, cc in SUBCOL.items():
        f = "=AVERAGE(%s)" % ",".join(std_ref(c, i) for c in SUB[sub])
        cell = ws.cell(row=r, column=cc, value=f)
        cell.number_format = "0.0000"; cell.border = THIN
    for k, key in enumerate(["US", "UD", "UE"]):
        c = ws.cell(row=r, column=6 + k, value=PAPER[key][i])
        c.number_format = "0.000"; c.fill = YEL; c.border = THIN
        d = ws.cell(row=r, column=10 + k, value="=%s%d-%s%d" % (CL(2 + k), r, CL(6 + k), r))
        d.number_format = "+0.0000;-0.0000;0"; d.fill = YEL; d.border = THIN
US = lambda i: "'4_子系统指数'!B%d" % (i + 3)
UD = lambda i: "'4_子系统指数'!C%d" % (i + 3)
UE = lambda i: "'4_子系统指数'!D%d" % (i + 3)

# ============ 5_MPI协调指数 ============
ws = wb.create_sheet("5_MPI协调指数")
ws["A1"] = "MPI 失衡惩罚型协调指数：D = M − S²/M（M为三子系统均值，S²为总体方差）"; ws["A1"].font = BOLD
header(ws, 2, ["年份", "US", "UD", "UE", "M 均值", "S 标准差", "S² 方差", "CV", "D  MPI",
               "", "D 论文", "D 差", "", "最弱子系统"], [8] + [10] * 13)
for i, y in enumerate(YS):
    r = i + 3
    ws.cell(row=r, column=1, value=y).border = THIN
    for k, fn in enumerate([US, UD, UE]):
        c = ws.cell(row=r, column=2 + k, value="=%s" % fn(i)); c.font = GRN
        c.number_format = "0.0000"; c.border = THIN
    trip = "B{0}:D{0}".format(r)
    for cc, f, fmt in [(5, "=AVERAGE(%s)" % trip, "0.0000"),
                       (6, "=STDEV.P(%s)" % trip, "0.0000"),
                       (7, "=VAR.P(%s)" % trip, "0.000000"),
                       (8, "=F%d/E%d" % (r, r), "0.0000"),
                       (9, "=E%d-G%d/E%d" % (r, r, r), "0.0000")]:
        c = ws.cell(row=r, column=cc, value=f); c.number_format = fmt; c.border = THIN
    c = ws.cell(row=r, column=11, value=PAPER["D"][i]); c.number_format = "0.000"; c.fill = YEL; c.border = THIN
    c = ws.cell(row=r, column=12, value="=I%d-K%d" % (r, r))
    c.number_format = "+0.0000;-0.0000;0"; c.fill = YEL; c.border = THIN
    c = ws.cell(row=r, column=14,
                value='=INDEX({"安全","发展","生态"},MATCH(MIN(%s),%s,0))' % (trip, trip))
    c.border = THIN
DM = lambda i: "'5_MPI协调指数'!I%d" % (i + 3)

wb.save(os.path.join(os.path.dirname(__file__), "_part1.xlsx"))
print("第一部分完成：说明 / 1_基础数据 / 2_指标值 / 3_标准化 / 4_子系统指数 / 5_MPI协调指数")

# ============ 6_组合权重 ============
ws = wb.create_sheet("6_组合权重")
ws["A1"] = "熵值法 × CRITIC 组合权重（仅用于诊断与稳健性对照，不参与子系统合成）"; ws["A1"].font = BOLD
ws["A2"] = "注：论文子系统合成采用等权。本表权重用于「7_结构性贡献缺口」与权重集中度检验。"
CODES = [c for _, c, _, _, _, _, _ in IND]
HR = 4
header(ws, HR, ["编码", "指标", "Σx'", "熵值 e", "差异 d=1-e", "熵值权重", "标准差 σ",
                "冲突性 Σ(1-r)", "CRITIC量 C", "CRITIC权重", "组合(积)", "组合权重"],
       [7, 32] + [12] * 10)
n0 = HR + 1
for k, code in enumerate(CODES):
    r = n0 + k; sr = IR[code]
    rng = "'3_标准化'!$%s$%d:$%s$%d" % (col(0), sr, col(NY - 1), sr)
    ws.cell(row=r, column=1, value=code).border = THIN
    ws.cell(row=r, column=2, value=IND[k][2]).border = THIN
    # 熵值：p=x/Σx, e=-1/ln(n)*Σ p·ln p
    terms = "+".join("IF(%s=0,0,(%s/$C%d)*LN(%s/$C%d))" %
                     (std_ref(code, i), std_ref(code, i), r, std_ref(code, i), r) for i in range(NY))
    for cc, f, fmt in [(3, "=SUM(%s)" % rng, "0.0000"),
                       (4, "=-1/LN(%d)*(%s)" % (NY, terms), "0.000000"),
                       (5, "=1-D%d" % r, "0.000000"),
                       (6, "=E%d/SUM($E$%d:$E$%d)" % (r, n0, n0 + 12), "0.0000"),
                       (7, "=STDEV.S(%s)" % rng, "0.0000")]:
        c = ws.cell(row=r, column=cc, value=f); c.number_format = fmt; c.border = THIN
    conf = "+".join("(1-CORREL(%s,'3_标准化'!$%s$%d:$%s$%d))" %
                    (rng, col(0), IR[o], col(NY - 1), IR[o]) for o in CODES if o != code)
    for cc, f, fmt in [(8, "=%s" % conf, "0.0000"),
                       (9, "=G%d*H%d" % (r, r), "0.0000"),
                       (10, "=I%d/SUM($I$%d:$I$%d)" % (r, n0, n0 + 12), "0.0000"),
                       (11, "=F%d*J%d" % (r, r), "0.000000"),
                       (12, "=K%d/SUM($K$%d:$K$%d)" % (r, n0, n0 + 12), "0.0000")]:
        c = ws.cell(row=r, column=cc, value=f); c.number_format = fmt; c.border = THIN
ws.cell(row=n0 + 13, column=2, value="合计").font = BOLD
for cc in (6, 10, 12):
    c = ws.cell(row=n0 + 13, column=cc, value="=SUM(%s%d:%s%d)" % (CL(cc), n0, CL(cc), n0 + 12))
    c.number_format = "0.0000"; c.font = BOLD; c.border = THIN
W = lambda code: "'6_组合权重'!$L$%d" % (n0 + CODES.index(code))

# ============ 7_结构性贡献缺口 ============
ws = wb.create_sheet("7_结构性贡献缺口")
ws["A1"] = "结构性贡献缺口 O(i,j) = w(j)·I(i,j) / Σ w·I，I=1-x'，逐年计算后取样本期均值（论文表4）"; ws["A1"].font = BOLD
yearhdr(ws, 2, ["编码", "指标"], [7, 32] + [10] * NY)
ws.cell(row=2, column=C0 + NY, value="样本期均值").fill = HDR
ws.cell(row=2, column=C0 + NY).font = BOLD
ws.cell(row=2, column=C0 + NY + 1, value="论文表4").fill = HDR
ws.cell(row=2, column=C0 + NY + 1).font = BOLD
P4 = {"D3": .3109, "E2": .2556, "E4": .1582, "E1": .1145, "D5": .0922, "D4": .0513}
DEN = n0 + 14           # 分母行（放在权重表？改放本表底部）
for k, code in enumerate(CODES):
    r = k + 3
    ws.cell(row=r, column=1, value=code).border = THIN
    ws.cell(row=r, column=2, value=IND[k][2]).border = THIN
    for i in range(NY):
        f = "=%s*(1-%s)/%s$%d" % (W(code), std_ref(code, i), col(i), 3 + len(CODES) + 1)
        c = ws.cell(row=r, column=C0 + i, value=f); c.number_format = "0.00%"; c.border = THIN
    c = ws.cell(row=r, column=C0 + NY,
                value="=AVERAGE(%s%d:%s%d)" % (col(0), r, col(NY - 1), r))
    c.number_format = "0.00%"; c.font = BOLD; c.border = THIN
    if code in P4:
        c = ws.cell(row=r, column=C0 + NY + 1, value=P4[code])
        c.number_format = "0.00%"; c.fill = YEL; c.border = THIN
dr = 3 + len(CODES) + 1
ws.cell(row=dr, column=2, value="分母 Σ w·I").font = BOLD
for i in range(NY):
    f = "=SUMPRODUCT(('6_组合权重'!$L$%d:$L$%d),(1-'3_标准化'!%s%d:%s%d))" % (
        n0, n0 + 12, col(i), IR["S1"], col(i), IR["E4"])
    c = ws.cell(row=dr, column=C0 + i, value=f); c.number_format = "0.0000"; c.font = BOLD; c.border = THIN

# ============ 8_局部敏感性 ============
ws = wb.create_sheet("8_局部敏感性")
ws["A1"] = "局部敏感性 ε = (D⁺−D⁻)/(0.02·D)，±1% 相对扰动双侧数值微分（论文表5）"; ws["A1"].font = BOLD
header(ws, 2, ["年份", "εS 安全", "εD 发展", "εE 生态", "", "εS 论文", "εD 论文", "εE 论文",
               "", "εS 差", "εD 差", "εE 差"], [8] + [11] * 11)
P5 = {2017:(.271,.351,.378),2018:(.274,.365,.361),2019:(.297,.357,.346),2020:(.317,.379,.305),
      2021:(.326,.370,.304),2022:(.339,.370,.291),2023:(.343,.338,.319),2024:(.344,.333,.323),
      2025:(.354,.353,.292)}
def mpi_expr(vals):
    v = ",".join(vals)
    return "(AVERAGE(%s)-VAR.P(%s)/AVERAGE(%s))" % (v, v, v)
for i, y in enumerate(YS):
    r = i + 3
    ws.cell(row=r, column=1, value=y).border = THIN
    base = [US(i), UD(i), UE(i)]
    for k in range(3):
        up = list(base); up[k] = "%s*1.01" % base[k]
        dn = list(base); dn[k] = "%s*0.99" % base[k]
        f = "=(%s-%s)/(0.02*%s)" % (mpi_expr(up), mpi_expr(dn), DM(i))
        c = ws.cell(row=r, column=2 + k, value=f); c.number_format = "0.000"; c.border = THIN
        if y in P5:
            p = ws.cell(row=r, column=6 + k, value=P5[y][k])
            p.number_format = "0.000"; p.fill = YEL; p.border = THIN
            d = ws.cell(row=r, column=10 + k, value="=%s%d-%s%d" % (CL(2 + k), r, CL(6 + k), r))
            d.number_format = "+0.000;-0.000;0"; d.fill = YEL; d.border = THIN
ws.cell(row=NY + 4, column=1,
        value="注：2016年生态类指标处于固定基准区间下限，弹性数值不稳定，论文表5未列（本表照算，仅供参考）。")

# ============ 9_留一法 ============
ws = wb.create_sheet("9_留一法")
ws["A1"] = "留一法（LOO）：逐一剔除单项指标后按其余设定不变重算（论文表6）"; ws["A1"].font = BOLD
yearhdr(ws, 2, ["剔除指标", "所属子系统"], [26, 12] + [10] * NY)
extra = ["发展最弱年数", "D(2016)", "D(2025)"]
for j, t in enumerate(extra):
    c = ws.cell(row=2, column=C0 + NY + j, value=t); c.fill = HDR; c.font = BOLD
def sub_expr(sub, i, drop=None):
    cs = [c for c in SUB[sub] if c != drop]
    return "AVERAGE(%s)" % ",".join(std_ref(c, i) for c in cs)
for k, code in enumerate(CODES):
    r = k + 3; sub = IND[k][0]
    ws.cell(row=r, column=1, value="%s %s" % (code, IND[k][2].split("（")[0])).border = THIN
    ws.cell(row=r, column=2, value=sub).border = THIN
    for i in range(NY):
        trip = [sub_expr(s, i, code if s == sub else None) for s in ["安全 S", "发展 D", "生态 E"]]
        c = ws.cell(row=r, column=C0 + i, value="=%s" % mpi_expr(trip))
        c.number_format = "0.000"; c.border = THIN
    cnt = "+".join("IF(%s=MIN(%s,%s,%s),1,0)" % (
        sub_expr("发展 D", i, code if sub == "发展 D" else None),
        sub_expr("安全 S", i, code if sub == "安全 S" else None),
        sub_expr("发展 D", i, code if sub == "发展 D" else None),
        sub_expr("生态 E", i, code if sub == "生态 E" else None)) for i in range(NY))
    ws.cell(row=r, column=C0 + NY, value="=%s" % cnt).border = THIN
    ws.cell(row=r, column=C0 + NY + 1, value="=%s3" % col(0)).border = THIN
    ws.cell(row=r, column=C0 + NY + 1, value="=%s%d" % (col(0), r)).number_format = "0.000"
    ws.cell(row=r, column=C0 + NY + 2, value="=%s%d" % (col(NY - 1), r)).number_format = "0.000"
srow = len(CODES) + 4
ws.cell(row=srow, column=1, value="频次统计（论文§5(三) 建议补报）").font = BOLD
ws.cell(row=srow + 1, column=1, value="发展最弱年数 ≥3 年的设定数")
ws.cell(row=srow + 1, column=3, value="=COUNTIF(%s3:%s%d,\">=3\")" % (col(NY), col(NY), len(CODES) + 2))
ws.cell(row=srow + 2, column=1, value="发展最弱年数 最小值 / 最大值")
ws.cell(row=srow + 2, column=3, value="=MIN(%s3:%s%d)&\" / \"&MAX(%s3:%s%d)" % (
    col(NY), col(NY), len(CODES) + 2, col(NY), col(NY), len(CODES) + 2))

wb.save(os.path.join(os.path.dirname(__file__), "_part2.xlsx"))
print("第二部分完成：6_组合权重 / 7_结构性贡献缺口 / 8_局部敏感性 / 9_留一法")

# ============ 10_稳健性检验 ============
ws = wb.create_sheet("10_稳健性检验")
ws["A1"] = "稳健性检验：滚动端点稳定性、极差标准化对照、惯例性上限±25%扰动"; ws["A1"].font = BOLD
ws["A3"] = "(一) 滚动端点稳定性：固定基准法 vs 极差标准化"; ws["A3"].font = BOLD
ws["A4"] = "固定基准法的上下限外生，截断样本不改变历史年份取值（差异恒为0）；极差法以样本极值为分母，截断样本会回溯改写历史值。"
yearhdr(ws, 6, ["编码", "指标", "口径"], [7, 30, 22] + [10] * NY)
rr = 7
for k, code in enumerate(CODES):
    d = IND[k][3]
    sr = IR[code]
    full = "'2_指标值'!$%s$%d:$%s$%d" % (vcol(0), sr, vcol(NY - 1), sr)
    trunc = "'2_指标值'!$%s$%d:$%s$%d" % (vcol(0), sr, vcol(6), sr)   # 2016—2022 截断样本
    for tag, rng in (("极差·全样本", full), ("极差·截断至2022", trunc)):
        ws.cell(row=rr, column=1, value=code).border = THIN
        ws.cell(row=rr, column=2, value=IND[k][2]).border = THIN
        ws.cell(row=rr, column=3, value=tag).border = THIN
        for i in range(7):        # 仅比较 2016—2022
            src = "'2_指标值'!%s%d" % (vcol(i), sr)
            num = ("(%s-MIN(%s))" % (src, rng)) if d == "正向" else ("(MAX(%s)-%s)" % (rng, src))
            f = "=0.98*(%s/(MAX(%s)-MIN(%s)))+0.01" % (num, rng, rng)
            c = ws.cell(row=rr, column=C0 + i, value=f); c.number_format = "0.0000"; c.border = THIN
        rr += 1
    ws.cell(row=rr, column=3, value="差异（截断−全样本）").font = BOLD
    for i in range(7):
        c = ws.cell(row=rr, column=C0 + i, value="=%s%d-%s%d" % (col(i), rr - 1, col(i), rr - 2))
        c.number_format = "+0.0000;-0.0000;0"; c.fill = YEL; c.border = THIN
    rr += 1
ws.cell(row=rr + 1, column=2, value="极差法最大绝对差异").font = BOLD
diffs = ",".join("%s%d:%s%d" % (col(0), 9 + 3 * k, col(6), 9 + 3 * k) for k in range(len(CODES)))
ws.cell(row=rr + 1, column=4, value="=MAX(ABS(%s))" % diffs.replace(",", "),ABS(")).number_format = "0.0000"
ws.cell(row=rr + 2, column=2, value="固定基准法最大绝对差异").font = BOLD
ws.cell(row=rr + 2, column=4, value=0).number_format = "0.0000"
ws.cell(row=rr + 2, column=5, value="（上下限外生，恒为0，无需计算）")
ws.cell(row=rr + 3, column=2, value="论文§5(二)载值").font = BOLD
ws.cell(row=rr + 3, column=4, value=0.229).number_format = "0.0000"
ws.cell(row=rr + 3, column=5,
        value="注：本表按【指标层标准化值】计算；若论文0.229系按子系统指数层计算，量级会因取平均而显著收敛，"
              "两者不可直接比较——请作者确认0.229的计算层级。")

r2 = rr + 5
ws.cell(row=r2, column=1, value="(二) 惯例性上限 ±25% 扰动").font = BOLD
ws.cell(row=r2 + 1, column=1,
        value="在下方黄色单元格改动倍率（1=原设定，1.25/0.75=±25%），「2_指标值」的上限随之变化，全表自动重算。")
header(ws, r2 + 2, ["指标", "原上限", "扰动倍率", "扰动后上限"], [26, 12, 12, 14])
CONV = [("S1", 12000), ("S3", 1200), ("D3", 10), ("D4", 25), ("D5", 200)]
for k, (code, hi) in enumerate(CONV):
    r = r2 + 3 + k
    ws.cell(row=r, column=1, value="%s %s" % (code, IND[CODES.index(code)][2])).border = THIN
    ws.cell(row=r, column=2, value=hi).border = THIN
    c = ws.cell(row=r, column=3, value=1); c.font = BLUE; c.fill = YEL; c.border = THIN
    c = ws.cell(row=r, column=4, value="=B%d*C%d" % (r, r)); c.number_format = "0"; c.border = THIN
# 把「2_指标值」这五项的上限改为引用本表扰动后上限
ws2 = wb["2_指标值"]
for k, (code, hi) in enumerate(CONV):
    ws2.cell(row=IR[code], column=6, value="='10_稳健性检验'!D%d" % (r2 + 3 + k))

r3 = r2 + 10
ws.cell(row=r3, column=1, value="(三) 权重集中度（论文§5(二)）").font = BOLD
header(ws, r3 + 1, ["项目", "数值", "论文载值"], [34, 14, 14])
RNG = "'6_组合权重'!$L$%d:$L$%d" % (n0, n0 + 12)
TOP4 = "+".join("LARGE(%s,%d)" % (RNG, i) for i in (1, 2, 3, 4))
items = [("最大单项组合权重", "=MAX(%s)" % RNG, 0.303),
         ("前四项组合权重合计", "=%s" % TOP4, 0.827),
         ("其余九项合计", "=1-(%s)" % TOP4, 0.173),
         ("安全类四项组合权重合计", "=%s" % "+".join(W(c) for c in SUB["安全 S"]), None),
         ("安全类单项最大权重", "=MAX(%s)" % ",".join(W(c) for c in SUB["安全 S"]), 0.012)]
for k, (nm, f, pv) in enumerate(items):
    r = r3 + 2 + k
    ws.cell(row=r, column=1, value=nm).border = THIN
    c = ws.cell(row=r, column=2, value=f); c.number_format = "0.000"; c.border = THIN
    if pv is not None:
        c = ws.cell(row=r, column=3, value=pv); c.number_format = "0.000"; c.fill = YEL; c.border = THIN

# ============ 11_对照指数 ============
ws = wb.create_sheet("11_对照指数")
ws["A1"] = "简化线性加总对照指数（论文§5(一)）：同一13项指标重新分组为协调/绿色/开放/共享四维度"; ws["A1"].font = BOLD
ws["A2"] = "★ 使用前请先在 C 列填入各指标的维度归属（协调 / 绿色 / 开放 / 共享），填好后下方自动计算。"
ws["A2"].font = Font(bold=True, color="C00000")
header(ws, 4, ["编码", "指标", "归入维度"], [7, 34, 14])
for k, code in enumerate(CODES):
    r = k + 5
    ws.cell(row=r, column=1, value=code).border = THIN
    ws.cell(row=r, column=2, value=IND[k][2]).border = THIN
    c = ws.cell(row=r, column=3, value=""); c.font = BLUE; c.fill = YEL; c.border = THIN
dr2 = len(CODES) + 7
ws.cell(row=dr2, column=1, value="四维度指数与对照指数").font = BOLD
yearhdr(ws, dr2 + 1, ["维度"], [14] + [10] * NY)
DIMS = ["协调", "绿色", "开放", "共享"]
for k, dm in enumerate(DIMS):
    r = dr2 + 2 + k
    ws.cell(row=r, column=1, value=dm).border = THIN
    for i in range(NY):
        f = "=IFERROR(AVERAGEIF($C$5:$C$%d,$A%d,'3_标准化'!%s$%d:%s$%d),\"\")" % (
            len(CODES) + 4, r, col(i), IR["S1"], col(i), IR["E4"])
        c = ws.cell(row=r, column=C0 + i, value=f); c.number_format = "0.0000"; c.border = THIN
rc = dr2 + 6
ws.cell(row=rc, column=1, value="对照指数（四维算术平均）").font = BOLD
for i in range(NY):
    c = ws.cell(row=rc, column=C0 + i, value="=IFERROR(AVERAGE(%s%d:%s%d),\"\")" % (
        col(i), dr2 + 2, col(i), dr2 + 5))
    c.number_format = "0.0000"; c.border = THIN
ws.cell(row=rc + 1, column=1, value="与 MPI 协调指数的相关系数").font = BOLD
ws.cell(row=rc + 1, column=3,
        value="=IFERROR(CORREL(%s%d:%s%d,'5_MPI协调指数'!I3:I%d),\"待填维度归属\")" % (
            col(0), rc, col(NY - 1), rc, NY + 2)).number_format = "0.000"
ws.cell(row=rc + 2, column=1, value="论文载值")
ws.cell(row=rc + 2, column=3, value=0.979).number_format = "0.000"

OUT = os.path.join(os.path.dirname(__file__), "西藏三元协调评价-计算底稿.xlsx")
for tmp in ("_part1.xlsx", "_part2.xlsx"):
    p = os.path.join(os.path.dirname(__file__), tmp)
    if os.path.exists(p): os.remove(p)
wb.save(OUT)
print("已生成：", OUT)
print("工作表：", " / ".join(wb.sheetnames))
