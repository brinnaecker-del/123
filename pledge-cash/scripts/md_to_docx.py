"""
把本项目的说明类 Markdown（标题、段落、列表、表格、加粗、链接）转换为 Word 文档。

    python md_to_docx.py 输入.md 输出.docx
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT_CN, FONT_EN = 'SimSun', 'Times New Roman'


def set_font(run, size=None, bold=None):
    run.font.name = FONT_EN
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT_CN)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_hyperlink(par, text, url, size):
    rid = par.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                             is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), rid)
    r = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), FONT_EN)
    fonts.set(qn('w:hAnsi'), FONT_EN)
    fonts.set(qn('w:eastAsia'), FONT_CN)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '1F4E79')
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(size * 2)))
    for e in (fonts, color, u, sz):
        rpr.append(e)
    r.append(rpr)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    h.append(r)
    par._p.append(h)


TOKEN = re.compile(r'(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\)|`[^`]+`)')


def add_inline(par, text, size=10.5, bold=False):
    text = text.replace('\\*', '*')
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            add_inline(par, part[2:-2], size, True)
        elif part.startswith('[') and '](' in part:
            label, url = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part).groups()
            add_hyperlink(par, label, url, size)
        elif part.startswith('`'):
            set_font(par.add_run(part[1:-1]), size, bold)
        else:
            set_font(par.add_run(part), size, bold)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcpr.append(shd)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
    header, body = cells[0], [r for r in cells[2:]]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate([header] + body):
        for j in range(len(header)):
            c = t.cell(i, j)
            c.text = ''
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, row[j] if j < len(row) else '', size=9.5, bold=(i == 0))
            if i == 0:
                shade(c, 'E7E6E6')
    doc.add_paragraph()


def convert(src, dst):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ('left_margin', 'right_margin'):
        setattr(sec, side, Cm(2.0))
    sec.top_margin = sec.bottom_margin = Cm(2.2)

    lines = open(src, encoding='utf-8').read().split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('|'):
            block = []
            while i < len(lines) and lines[i].startswith('|'):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            p = doc.add_paragraph()
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, line.lstrip('#').strip(), size={1: 16, 2: 13, 3: 11.5}.get(level, 11), bold=True)
            p.paragraph_format.space_before = Pt(10 if level > 1 else 0)
            p.paragraph_format.space_after = Pt(6)
        elif re.match(r'\s*[-*] ', line):
            p = doc.add_paragraph()
            indent = (len(line) - len(line.lstrip())) // 2
            p.paragraph_format.left_indent = Cm(0.6 + 0.6 * indent)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            add_inline(p, '• ' + re.sub(r'^\s*[-*] ', '', line))
        elif line.startswith('>'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, line.lstrip('> ').strip())
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        elif line.strip():
            p = doc.add_paragraph()
            add_inline(p, line.strip())
        i += 1
    doc.save(dst)


if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
    print('已写出', sys.argv[2])
