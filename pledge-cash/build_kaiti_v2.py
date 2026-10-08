"""生成开题报告修订稿：在原稿副本上替换与数据有关的内容、插入进一步检验的结果。原稿不改动。

用法：python -I build_kaiti_v2.py
所有数值从 results/*.json 读取；“显著/不显著”等措辞由显著性星号自动生成。
"""
import copy, json, os, re
import docx
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
R = lambda f: json.load(open(os.path.join(HERE, 'results', f), encoding='utf-8'))
BL, CH, BD, DD, RB, MC = (R('res_baseline.json'), R('res_channels.json'), R('res_boundary.json'),
                          R('res_did.json'), R('res_robust.json'), R('res_measure_check.json'))
SRC = os.path.join(HERE, '开题报告-原稿.docx')
OUT = os.path.join(HERE, '开题报告-修订稿.docx')
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

# ---------------- 数值格式 ----------------
def stars(c):
    return len(re.findall(r'\*', c[0]))
def sig(c):
    n = stars(c)
    return {3: '在1%水平上显著', 2: '在5%水平上显著', 1: '在10%水平上显著', 0: '不显著'}[n]
def cs(c):  # 系数(标准误)
    return f'{c[0]}{c[1]}'
def need(cond, msg):
    if not cond:
        raise AssertionError('结果与正文表述不符：' + msg)
ns = lambda c: stars(c) == 0          # 不显著
neg = lambda c: c[0].lstrip().startswith('-')
f3 = lambda x: f'{x:.3f}'
f4 = lambda x: f'{x:.4f}'
pct = lambda x, d=1: f'{100 * x:.{d}f}%'
num = lambda x: f'{int(x):,}'

# ---------------- XML 工具 ----------------
def ptxt(p):
    return ''.join(t.text or '' for t in p.iter(W + 't'))

def set_text(p, text):
    runs = [r for r in p.iter(W + 'r')]
    assert runs, 'paragraph has no run'
    withtext = [r for r in runs if r.find(W + 't') is not None]
    r0 = withtext[0] if withtext else runs[0]
    for r in runs:
        if r is not r0 and r.getparent() is not None:
            r.getparent().remove(r)
    for sub in list(r0):
        if sub.tag not in (W + 'rPr', W + 't'):
            r0.remove(sub)
    ts = r0.findall(W + 't')
    for t in ts[1:]:
        r0.remove(t)
    t = ts[0] if ts else etree.SubElement(r0, W + 't')
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return p

def cell_children(cell):
    return list(cell._tc.iterchildren())

def find_p(cell, pred):
    hits = [ch for ch in cell._tc.iterchildren() if ch.tag == W + 'p' and pred(ptxt(ch))]
    assert len(hits) == 1, (len(hits), [ptxt(h)[:30] for h in hits])
    return hits[0]

def replace_in(p, old, new):
    t = ptxt(p)
    assert old in t, (old, t[:80])
    set_text(p, t.replace(old, new))

def clone(ref, text):
    e = copy.deepcopy(ref)
    if text is not None:
        set_text(e, text)
    return e

def set_tc(tc, text):
    ps = tc.findall(W + 'p')
    for p in ps[1:]:
        tc.remove(p)
    set_text(ps[0], text)

def row_cells(tr):
    return tr.findall(W + 'tc')

# 新表：与原稿表格同格式（全框线、表头灰底加粗、五号字、Times New Roman/宋体）
def make_table(rows, widths, header_rows=1, first_left=True):
    total = sum(widths)
    tbl = etree.SubElement(etree.Element('dummy'), W + 'tbl')
    tblPr = etree.SubElement(tbl, W + 'tblPr')
    etree.SubElement(tblPr, W + 'tblW', {W + 'type': 'dxa', W + 'w': str(total)})
    etree.SubElement(tblPr, W + 'jc', {W + 'val': 'center'})
    grid = etree.SubElement(tbl, W + 'tblGrid')
    for w in widths:
        etree.SubElement(grid, W + 'gridCol', {W + 'w': str(w)})
    for ri, row in enumerate(rows):
        tr = etree.SubElement(tbl, W + 'tr')
        trPr = etree.SubElement(tr, W + 'trPr')
        etree.SubElement(trPr, W + 'cantSplit')
        if ri < header_rows:
            etree.SubElement(trPr, W + 'tblHeader')
        for ci, (txt, w) in enumerate(zip(row, widths)):
            tc = etree.SubElement(tr, W + 'tc')
            tcPr = etree.SubElement(tc, W + 'tcPr')
            etree.SubElement(tcPr, W + 'tcW', {W + 'type': 'dxa', W + 'w': str(w)})
            b = etree.SubElement(tcPr, W + 'tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                etree.SubElement(b, W + side, {W + 'val': 'single', W + 'color': '000000', W + 'sz': '4'})
            if ri < header_rows:
                etree.SubElement(tcPr, W + 'shd', {W + 'fill': 'F2F2F2', W + 'val': 'clear'})
            mar = etree.SubElement(tcPr, W + 'tcMar')
            for side, v in (('top', 60), ('left', 80), ('bottom', 60), ('right', 80)):
                etree.SubElement(mar, W + side, {W + 'type': 'dxa', W + 'w': str(v)})
            etree.SubElement(tcPr, W + 'vAlign', {W + 'val': 'center'})
            p = etree.SubElement(tc, W + 'p')
            pPr = etree.SubElement(p, W + 'pPr')
            etree.SubElement(pPr, W + 'spacing', {W + 'after': '0', W + 'line': '280'})
            etree.SubElement(pPr, W + 'jc', {W + 'val': 'left' if (ci == 0 and first_left) else 'center'})
            r = etree.SubElement(p, W + 'r')
            rPr = etree.SubElement(r, W + 'rPr')
            etree.SubElement(rPr, W + 'rFonts', {W + 'ascii': 'Times New Roman', W + 'cs': 'Times New Roman',
                                                 W + 'eastAsia': 'SimSun', W + 'hAnsi': 'Times New Roman'})
            if ri < header_rows:
                etree.SubElement(rPr, W + 'b'); etree.SubElement(rPr, W + 'bCs')
            etree.SubElement(rPr, W + 'sz', {W + 'val': '21'}); etree.SubElement(rPr, W + 'szCs', {W + 'val': '21'})
            t = etree.SubElement(r, W + 't'); t.text = str(txt)
            t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return tbl

doc = docx.Document(SRC)
T = doc.tables
c1, c2, c3, c4, c5 = (T[i].cell(0, 0) for i in (1, 2, 3, 4, 5))

# ======================= 一、选题依据 =======================
p = find_p(c1, lambda t: t.startswith('在现实层面'))
replace_in(p, '本文拟进一步检验质押是否通过提高现金储备而抑制企业资本支出与创新投入', '本文拟进一步检验质押对企业资本支出与创新投入的影响')
p = find_p(c1, lambda t: t.startswith('综上，现有文献'))
replace_in(p, '并进一步检验质押是否通过现金储备渠道挤出企业投资支出', '并进一步检验质押对企业投资支出的影响')

# ======================= 二、研究内容 =======================
p = find_p(c2, lambda t: t.startswith('第五章，实证结果分析'))
replace_in(p, '投资挤出的拓展检验', '质押与企业投资的拓展检验')
p = find_p(c2, lambda t: t.startswith('在实证层面'))
replace_in(p, '同时检验质押是否通过现金储备渠道挤出企业投资支出', '同时检验质押对企业投资支出的影响')

# ======================= 三、困难与对策 =======================
sd = BD['sa_diag']
p = find_p(c3, lambda t: t.startswith('第二，融资约束指标的测度误差问题'))
replace_in(p, '本文拟按原文做法截尾并保留原始符号，以SA越大表示融资约束越强，同时报告SA与企业规模、上市年限的相关系数以核验方向；',
           f'需要注意，原文系数以2004年美元计，若直接以人民币百万元计算企业规模，二次项的拐点将降至约{sd["turning_point_yi_yuan_rmb"]:.1f}亿元，'
           f'样本中约{pct(sd["share_above_tp_rmb"], 0)}的观测会落入规模效应反转的区间。本文因此先按2004年汇率（8.28元/美元）将总资产折算为百万美元，'
           f'再按原文做法截尾并保留原始符号，以SA越大表示融资约束越强。经核验，SA与企业规模、上市年限的相关系数分别为{sd["corr_SA_Size"]:.2f}和{sd["corr_SA_Age"]:.2f}，方向符合预期；')
p = find_p(c3, lambda t: t.startswith('第四，被解释变量的计量可靠性问题'))
replace_in(p, '投向上市公司的记录仅约占2%', '投向上市公司的记录仅约占1.7%')

# ======================= 四、已具备的条件 =======================
# 表1：明细表没有预警线/平仓线字段
tbls4 = [ch for ch in c4._tc.iterchildren() if ch.tag == W + 'tbl']
t1, t2, t3, t4, t5, t6, t7 = tbls4
tc = row_cells(t1.findall(W + 'tr')[1])[0]
set_tc(tc, ptxt(tc).replace('、预警线／平仓线', ''))

p = find_p(c4, lambda t: t.startswith('第 1 步'))
set_text(p, '第 1 步：对“股东股权质押统计表”中每一出质方按变动日期累加“数量增减”，在“全部解押”且“质押笔数”为空时归零'
            '（转增股本、部分解押记录的“质押笔数”同样为空，但质押并未解除，不予归零），得到任意时点的未解押质押股数；'
            '取每年12月31日（含）前的最后一条记录作为年末余额，按股东名称（含该出质方在统计表中使用过的全部名称）匹配“十大股东文件”中的第一大股东，'
            '除以其持股数得到Pledge（统计表存在漏记解押，累加结果大于1或小于0的截取至[0,1]），并生成Pledge_Dum与占总股本口径的Pledge_Ratio2；'
            '另以“股东股权质押情况明细表”逐笔剩余质押数量加总构造稳健性口径；')
p = find_p(c4, lambda t: t.startswith('第 3 步'))
set_text(p, '第 3 步：以“股票代码＋会计年度”为主键 merge 质押数据与财务、治理、市场数据，剔除金融业、ST／*ST（依据当年年报证券简称识别）、'
            '资不抵债及主要变量缺失的样本；股权性质缺失的观测予以保留（基准回归不使用产权性质），仅在产权性质相关检验中剔除；')
p = find_p(c4, lambda t: t.startswith('第 4 步'))
set_text(p, '第 4 步：计算 Cash、SA（及稳健性替代的 KZ、WW）以及各控制变量，对连续变量在 1%／99% 分位上缩尾（winsorize），'
            '分位数取自剔除金融、ST、资不抵债及 Cash、Growth、TobinQ 缺失后的样本；Top1 与质押比例类变量不缩尾；')
pl_pos = BL['t4']['Pledge_Dum']['mean']
p = find_p(c4, lambda t: t.startswith('截至目前'))
set_text(p, f'截至目前，第 1—6 步均已完成，基准回归、U型关系复现、渠道检验、边界条件、2018年股票质押新规双重差分及稳健性与内生性检验已得到初步结果'
            f'（见本部分（三））；现金边际价值、KZ／WW 指数、被解释变量可靠性、投资效率与纾困政策等检验所需数据尚待补充下载（见本部分（四））。'
            f'现已形成2007—2025年沪深A股非金融上市公司的公司—年度面板，样本共{num(BL["N"])}个观测、{num(BL["firms"])}家公司。'
            f'质押比例的构造经过三项交叉验证：统计表口径与明细表口径的相关系数为{MC["corr_stat_detail"]:.2f}，{pct(MC["share_within_5pp"], 0)}的公司—年度差异在5个百分点以内；'
            f'由此得到的“是否质押”与十大股东文件中的质押标识一致率为{pct(MC["agree_dum_flag1"])}。'
            f'另有{pct(MC["share_raw_gt1"])}的观测累加结果超过1（统计表漏记解押），已截取至1；'
            f'{pct(MC["share_stale_among_pledgers"])}的质押观测仅由三年以前的统计表记录支撑，其影响在稳健性检验中报告。')

# 表2：变量定义
def t2row(sym):
    for tr in t2.findall(W + 'tr'):
        tcs = row_cells(tr)
        if len(tcs) >= 3 and ptxt(tcs[2]).strip() == sym:
            return tcs
    raise KeyError(sym)
set_tc(t2row('Pledge')[3], '大股东质押股数/其持股总数，截取至[0,1]')
set_tc(t2row('Pledge_Ratio2')[3], '大股东质押股数/公司总股本（质押股数以其持股数为上限）')
set_tc(t2row('SA')[3], 'SA=−0.737×Size+0.043×Size²−0.040×Age（Size为总资产按2004年汇率折算的百万美元对数，上限4,500；Age为上市年数，上限37年；'
                       '保留原始符号，数值越大融资约束越强）；检验H3时以滞后一期SA在年度内分组，DID部分以2016—2017年均值固定分组（唐玮等，2019）')
set_tc(t2row('OREC')[3], '（其他应收款－其他应付款）/总资产，再减去同行业（制造业按两位大类、其余按门类）同年度均值')
set_tc(t2row('Pressure')[3], '对第一大股东每笔年末未解押的质押，以起始日收盘价为基期、按考虑现金红利再投资的月回报率推算复权价格，'
                             '判断当年是否跌破预警价格（明细表未提供预警线字段，按质押率40%、年利率10%、预警线160%设定）；'
                             '取触及预警线的质押股数占其全部未解押质押股数的比例')
set_tc(t2row('SOE')[3], '股权性质含“国企”=1，其余=0；股权性质缺失的观测在产权性质相关检验中剔除；多数公司不随时间变化，被公司固定效应吸收，主要用于H4分组')
set_tc(t2row('PledgeUse')[3], '质押资金投向上市公司=1，投向控股股东自身或第三方=0；投向上市公司的记录约占1.7%，仅作补充检验')
for tr in t2.findall(W + 'tr'):
    tcs = row_cells(tr)
    if len(tcs) >= 4 and ptxt(tcs[2]).strip() == 'Ind FE':
        set_tc(tcs[3], '制造业按证监会2012版两位大类、其余按门类设定（稳健性检验中与公司固定效应互换）')
# 表3：SA
for tr in t3.findall(W + 'tr'):
    tcs = row_cells(tr)
    if ptxt(tcs[0]).strip().startswith('SA'):
        set_tc(tcs[2], '系数 -0.737／0.043／-0.040；Size 为总资产按2004年汇率（8.28元/美元）折算的百万美元的自然对数，Age 为上市年限；'
                       '按原文做法对 Size、Age 截尾（上限分别为45亿美元与37年）；保留原始符号，数值越大表示融资约束越严重')

# 模型说明段：表述与实际一致（Controls 名单不变）
p = find_p(c4, lambda t: t.startswith('其中，i 表示公司'))
assert 'HighFC 为滞后一期SA指数高于当年中位数的虚拟变量' in ptxt(p)

# 表4：描述性统计（Cash2 所需交易性金融资产数据未提供，删去该行）
t4rows = t4.findall(W + 'tr')
VMAP = {}
for tr in t4rows[1:]:
    VMAP[ptxt(row_cells(tr)[0]).strip()] = tr
t4.remove(VMAP.pop('Cash2'))
for v, tr in VMAP.items():
    s = BL['t4'][v]
    vals = [num(s['count']), f3(s['mean']), f3(s['std']), f3(s['min']), f3(s['50%']), f3(s['max'])]
    for tc, x in zip(row_cells(tr)[1:], vals):
        set_tc(tc, x)
p = find_p(c4, lambda t: t.startswith('注：样本为2007—2025年沪深A股非金融上市公司'))
set_text(p, f'注：样本为2007—2025年沪深A股非金融上市公司，共{num(BL["N"])}个公司—年度观测、{num(BL["firms"])}家公司；已剔除ST、*ST公司；'
            f'连续变量经上下1%缩尾（Top1与质押比例不缩尾，质押比例截取至[0,1]）；SOE的观测值少于其他变量，系剔除股权性质缺失的观测所致；'
            f'SA指数按Hadlock和Pierce（2010）原文的美元单位计算，保留原始符号，数值越大表示融资约束越强。'
            f'现金持有的替代口径Cash2所需的交易性金融资产数据尚待补充，暂未列示。')

# 初步回归结果文字
t6r = BL['t6']
p = find_p(c4, lambda t: t.startswith('表6报告了模型1的初步估计结果'))
set_text(p, f'表6报告了模型1的初步估计结果。在控制公司与年度固定效应后，大股东股权质押比例的系数在各列中均显著为负：'
            f'列(2)加入控制变量后系数为{t6r[1]["coef"][0].rstrip("*")}，{sig(t6r[1]["coef"])}；列(3)以行业×年度固定效应吸收行业层面的共同冲击，'
            f'列(4)以上市年数固定效应控制新上市公司募集资金所带来的现金持有生命周期特征，系数分别为{t6r[2]["coef"][0]}与{t6r[3]["coef"][0]}，均保持显著为负。'
            f'按样本标准差折算，质押比例上升一个标准差，现金持有水平约下降其均值的{pct(-BL["t6_1sd_over_mean"])}。'
            f'这一净效应的符号与H1b（掏空说）的预测一致，但净效应只是两种机制的合力，掏空机制本身是否存在须由渠道检验判断（见下文表8）。')
t6rows = t6.findall(W + 'tr')
for j, c in enumerate(t6r):
    set_tc(row_cells(t6rows[1])[j + 1], c['coef'][0])
    set_tc(row_cells(t6rows[2])[j + 1], c['coef'][1])
for tr in t6rows:
    lab = ptxt(row_cells(tr)[0]).strip()
    if lab == '调整R²':
        for j, c in enumerate(t6r): set_tc(row_cells(tr)[j + 1], c['r2'])
    if lab == '观测值':
        for j, c in enumerate(t6r): set_tc(row_cells(tr)[j + 1], c['N'])
p = find_p(c4, lambda t: t.startswith('注：被解释变量为Cash；括号内为公司层面聚类稳健标准误'))
set_text(p, '注：被解释变量为Cash；括号内为公司层面聚类稳健标准误；***、**、*分别表示1%、5%、10%显著水平；连续变量经上下1%缩尾；'
            '行业×年度固定效应中，制造业按证监会2012版两位大类、其余按门类划分；观测值为剔除单例固定效应后的有效样本。下表同。')

t7r = BL['t7']
p = find_p(c4, lambda t: t.startswith('李常青等（2018b）基于2013—2015年季度数据发现'))
set_text(p, f'李常青等（2018b）基于2013—2015年季度数据发现，控股股东股权质押比例与现金持有呈U型关系。为便于比较，表7列(1)、(2)沿用其混合回归口径，'
            f'列(3)、(4)进一步控制公司固定效应。结果显示，按原文口径可以复现U型关系，拐点约为{t7r[0]["tp"]:.2f}，与原文报告的约55%基本一致；'
            f'但控制公司固定效应后，平方项在2013—2015年样本中不再显著，全样本中U型右端的斜率亦不显著，U型检验不成立（Lind和Mehlum，2010）。'
            f'这意味着既有文献发现的U型上升段可能主要反映高质押公司与其他公司之间的固有差异，而非同一公司质押加深后增持现金的行为。')
t7rows = t7.findall(W + 'tr')
lab_rows = {}
for i, tr in enumerate(t7rows):
    lab_rows.setdefault(ptxt(row_cells(tr)[0]).strip(), []).append(tr)
for j, c in enumerate(t7r):
    set_tc(row_cells(lab_rows['Pledge'][0])[j + 1], c['b1'][0])
    set_tc(row_cells(t7rows[t7rows.index(lab_rows['Pledge'][0]) + 1])[j + 1], c['b1'][1])
    set_tc(row_cells(lab_rows['Pledge²'][0])[j + 1], c['b2'][0])
    set_tc(row_cells(t7rows[t7rows.index(lab_rows['Pledge²'][0]) + 1])[j + 1], c['b2'][1])
    set_tc(row_cells(lab_rows['拐点'][0])[j + 1], f'{c["tp"]:.2f}' if c['holds'] else '—')
    set_tc(row_cells(lab_rows['U型检验'][0])[j + 1], '成立' if c['holds'] else '不成立')
    set_tc(row_cells(lab_rows['调整R²'][0])[j + 1], c['r2'])
    set_tc(row_cells(lab_rows['观测值'][0])[j + 1], c['N'])
p = find_p(c4, lambda t: t.startswith('注：列(1)、(2)按李常青等（2018b）的口径'))
set_text(p, '注：列(1)、(2)按李常青等（2018b）的口径，采用混合OLS并控制行业（制造业按两位大类、其余按门类）与年度虚拟变量；各列均剔除上市当年观测；'
            'U型检验采用Lind和Mehlum（2010）的端点斜率检验，质押比例为0与1处的斜率分别显著为负、显著为正时判定U型成立；未成立时不报告拐点。')

# ======================= 四(三)续：进一步检验 =======================
anchor = find_p(c4, lambda t: t.startswith('上述结果仅为初步估计'))
P_SUB = find_p(c4, lambda t: t == '初步回归结果')
P_BODY = find_p(c4, lambda t: t.startswith('表6报告了模型1的初步估计结果'))
P_CAP = find_p(c4, lambda t: t.startswith('表6　'))
P_NOTE = find_p(c4, lambda t: t.startswith('注：被解释变量为Cash；括号内为公司层面聚类稳健标准误'))
P_SEC = find_p(c4, lambda t: t.startswith('（三）实证设计与初步回归结果'))
P_ITEM = find_p(c4, lambda t: t.startswith('① 大股东股权质押数据'))
P_FIGCAP = find_p(c2, lambda t: t.startswith('图2　'))
cur = anchor
def add(el):
    global cur
    cur.addnext(el); cur = el
    return el
def body(t): return add(clone(P_BODY, t))
def sub(t): return add(clone(P_SUB, t))
def cap(t): return add(clone(P_CAP, t))
def note(t): return add(clone(P_NOTE, t))
def table(rows, widths, h=1): return add(make_table(rows, widths, h))
WD5 = [2400, 1320, 1320, 1320, 1320, 1320]
WD4 = [2400, 1650, 1650, 1650, 1650]
yes = lambda n: ['是'] * n

newsub = clone(P_SUB, '进一步检验的初步结果')
anchor.getparent().replace(anchor, newsub)
cur = newsub

body('在基准结果的基础上，本文已利用现有数据完成渠道检验、边界条件、2018年股票质押新规双重差分以及稳健性与内生性检验。'
     '各项检验的判定规则均在估计之前依据第二部分的假设写定，以下按规则如实报告；尚缺数据、未能估计的检验见本部分（四）。')

# ---- 1. 渠道检验 ----
c5_ = CH['t5_5']; c6_ = CH['t5_6']; pdg = CH['pressure_diag']
need(all(ns(c['PxP']) for c in c5_), '表8交乘项均不显著')
need(all(ns(c['PxP']) for c in CH['t5_5_measured_only']), '剔除无可测度事件后交乘项仍不显著')
need(ns(CH['t5_5_asof'][1]['PxP']), '年末已知口径交乘项不显著')
need(all(ns(c['Pledge']) for c in c6_), 'OREC 回归不显著')
need(all(not ns(c['Pledge']) and neg(c['Pledge']) for c in c5_), '加入平仓压力后 Pledge 仍显著为负')
sub('1．渠道检验：平仓压力与控股股东资金占用')
body(f'平仓压力Pressure按表2的定义构造。为避免明细表中已实际解押但未披露解押记录的质押被计入，剔除以事件全部记录判断、合同结束日早于当年末的质押，'
     f'以及无结束日期且起始日早于当年末3年以上的质押；当年无可测度质押的公司—年度取0。起始月内发生送转股的，按当月除权因子调整基期价格。'
     f'在有质押的公司—年度中，Pressure大于0的占{pct(pdg["Pressure"]["share_pos_among_pledgers"])}。为剥离股价下跌本身对现金的直接影响，模型控制当年股票回报率Ret，'
     f'并构造市场驱动的平仓压力Pressure_Mkt，即以沪深A股按上月总市值加权的月回报构造的市场指数替代个股价格路径（Pressure_Mkt大于0的占{pct(pdg["Pressure_Mkt"]["share_pos_among_pledgers"])}）。')
body(f'表8列(1)—(4)中，质押与平仓压力的交乘项均不显著：个股口径为{cs(c5_[0]["PxP"])}，滞后一期为{cs(c5_[1]["PxP"])}，'
     f'市场驱动口径为{cs(c5_[2]["PxP"])}，加入行业×年度固定效应后为{cs(c5_[3]["PxP"])}；Pledge的系数则保持显著为负。按判定规则，H2a未获支持。'
     f'稳健性方面，剔除有质押但无可测度事件的公司—年度后，交乘项仍不显著（市场驱动口径为{cs(CH["t5_5_measured_only"][2]["PxP"])}），'
     f'但该子样本中Pledge的系数降为{cs(CH["t5_5_measured_only"][0]["Pledge"])}；'
     f'若陈旧记录的判断只使用当年末之前已披露的结束日期，市场驱动口径的交乘项为{cs(CH["t5_5_asof"][1]["PxP"])}，同样不显著。'
     f'列(5)以控股股东资金净占用OREC为被解释变量，Pledge的系数为{cs(c6_[0]["Pledge"])}，{sig(c6_[0]["Pledge"])}；未经行业调整的口径为{cs(c6_[1]["Pledge"])}，'
     f'加入行业×年度固定效应后为{cs(c6_[2]["Pledge"])}。即在公司固定效应设定下，质押比例上升并未伴随以其他应收款衡量的资金占用增加，H2b中资金占用的部分未获支持。')
cap('表8　渠道检验：平仓压力与控股股东资金占用')
rows = [['变量', '(1)', '(2)', '(3)', '(4)', '(5)'],
        ['被解释变量', 'Cash', 'Cash', 'Cash', 'Cash', 'OREC'],
        ['压力口径', 'Pressure', 'Pressure滞后一期', 'Pressure_Mkt', 'Pressure_Mkt', '—']]
for lab, key in [('Pledge', 'Pledge'), ('Pledge×压力', 'PxP'), ('压力', 'P'), ('Ret', 'Ret')]:
    rows.append([lab] + [c[key][0] for c in c5_] + [c6_[0]['Pledge'][0] if key == 'Pledge' else ''])
    rows.append([''] + [c[key][1] for c in c5_] + [c6_[0]['Pledge'][1] if key == 'Pledge' else ''])
rows += [['控制变量'] + yes(5), ['公司固定效应'] + yes(5), ['年度固定效应', '是', '是', '是', '否', '是'],
         ['行业×年度固定效应', '否', '否', '否', '是', '否'],
         ['调整R²'] + [c['r2'] for c in c5_] + [c6_[0]['r2']], ['观测值'] + [c['N'] for c in c5_] + [c6_[0]['N']]]
table(rows, WD5, 3)
note('注：括号内为公司层面聚类稳健标准误；***、**、*分别表示1%、5%、10%显著水平。现金边际价值检验（模型4）所需的非经常性损益、利息费用、研发支出、'
     '现金股利与筹资活动现金流数据尚待补充，暂未估计。')

# ---- 2. 边界条件 ----
b8 = BD['t5_8']; d9 = BD['t5_9']; b10 = BD['t5_10']
need(not ns(b8[0]['PxH']) and neg(b8[0]['PxH']) and neg(b8[0]['Pledge']), '列(1)交乘项与 Pledge 同号且显著')
need(ns(b8[1]['PxH']) and not ns(b8[1]['PxAge']), '列(2)交乘项不显著、Pledge×Age 显著')
need(ns(b8[3]['PxH']), '人民币口径结论相同')
need(ns(b10[0]['PxSOE']) and ns(b10[3]['PxSOE']) and ns(b10[1]['Pledge']) and ns(b10[2]['Pledge']), 'H4 各项不显著')
sub('2．边界条件：融资约束与产权性质')
body(f'按模型(2)，以滞后一期SA指数高于当年中位数定义HighFC，并对质押与调节变量中心化后构造交乘项。表9列(1)中Pledge×HighFC为{cs(b8[0]["PxH"])}，'
     f'与Pledge同号且{sig(b8[0]["PxH"])}，单看支持H3。但SA指数由规模与上市年限构成，HighFC与上市年限的相关系数为{BD["corr_HighFC_Age"]:.2f}，'
     f'列(2)加入质押与中心化的上市年限、企业规模的交乘项后，Pledge×HighFC变为{cs(b8[1]["PxH"])}，不再显著，而Pledge×Age为{cs(b8[1]["PxAge"])}，{sig(b8[1]["PxAge"])}。'
     f'这说明列(1)中的调节作用实际来自上市年限：上市时间越短，质押与现金持有的负向关系越强。按事先写定的规则，H3不成立。'
     f'若沿用人民币单位计算SA，结论相同（加入交乘项后Pledge×HighFC为{cs(b8[3]["PxH"])}）。')
body(f'产权性质方面，剔除股权性质缺失的{num(BD["soe_coding"]["n_missing_excluded"])}个观测后，国有企业中存在质押的观测占{pct(d9[0]["dum"])}，'
     f'非国有企业为{pct(d9[1]["dum"])}；国有企业组质押比例的组内标准差为{d9[0]["within_sd"]}，低于非国有企业的{d9[1]["within_sd"]}。'
     f'表9列(3)中Pledge×SOE为{cs(b10[0]["PxSOE"])}，{sig(b10[0]["PxSOE"])}，H4未获支持；允许国有与非国有企业有各自的年度效应后（列(4)），'
     f'交乘项为{cs(b10[3]["PxSOE"])}，结论不变。分组回归中，非国有企业组与国有企业组的系数分别为{cs(b10[1]["Pledge"])}与{cs(b10[2]["Pledge"])}，均不显著。')
cap('表9　边界条件：融资约束与产权性质')
rows = [['变量', '(1)', '(2)', '(3)', '(4)'], ['调节变量', 'HighFC', 'HighFC', 'SOE', 'SOE']]
def cc(c, k):
    v = c.get(k)
    return v if v and v[0] else ['', '']
specs = [b8[0], b8[1], b10[0], b10[3]]
for lab, key in [('Pledge', 'Pledge'), ('Pledge×HighFC', 'PxH'), ('HighFC', 'HighFC'), ('Pledge×Age', 'PxAge'),
                 ('Pledge×Size', 'PxSize'), ('Pledge×SOE', 'PxSOE')]:
    vals = [cc(c, key) for c in specs]
    rows.append([lab] + [v[0] for v in vals]); rows.append([''] + [v[1] for v in vals])
rows += [['控制变量'] + yes(4), ['公司固定效应'] + yes(4), ['年度固定效应', '是', '是', '是', '否'],
         ['产权性质×年度固定效应', '否', '否', '否', '是'],
         ['调整R²'] + [c['r2'] for c in specs], ['观测值'] + [c['N'] for c in specs]]
table(rows, WD4, 2)
note('注：交乘项中的Pledge、HighFC、Age、Size与SOE均已中心化，主效应为调节变量取样本均值时的效应；列(3)、(4)剔除股权性质缺失的观测，SOE主效应在列(4)中被吸收。'
     'KZ指数与WW指数所需的现金股利、长期借款与应付债券数据尚待补充。')

# ---- 3. 2018 新规 ----
ds = DD['did_sample']; f11 = DD['t5_11']; m12 = DD['t5_12']; mr = DD['mean_reversion']; pp = DD['placebo_perm']
need(not ns(f11[0]['DID']) and neg(f11[0]['DID']), '第一阶段质押比例显著下降')
need(not ns(f11[1]['DID']) and not neg(f11[1]['DID']), '第一阶段平仓压力显著上升')
need(ns(f11[2]['DID']), '第一阶段资金占用不显著')
need(all(ns(c['coef']) for c in m12), '新规对现金持有各列均不显著')
sub('3．政策评估：2018年股票质押新规')
body(f'双重差分样本为2016年之前上市、且2016—2017年有质押数据的{num(ds["firms"])}家公司（{num(ds["N"])}个观测），处理强度PledgePre为2016—2017年大股东平均质押比例，'
     f'均值为{f3(ds["PledgePre_mean"])}、中位数为{f3(ds["PledgePre_median"])}。表10列(1)—(3)为“第一阶段”：新规实施后，处理强度越高的公司质押比例相对下降'
     f'（{cs(f11[0]["DID"])}），但这一下降并未超过均值回归的幅度——在窗口结构相同的设定下，真实队列与以2012—2013年为处理强度、2014年为虚拟政策时点的安慰剂队列的系数分别为'
     f'{mr["short"]["main"]["DID"][0]}与{mr["short"]["placebo"]["DID"][0]}（{mr["short"]["main"]["window"]}对{mr["short"]["placebo"]["window"]}），'
     f'{mr["medium"]["main"]["DID"][0]}与{mr["medium"]["placebo"]["DID"][0]}（{mr["medium"]["main"]["window"]}对{mr["medium"]["placebo"]["window"]}）。'
     f'平仓压力则相对上升（{cs(f11[1]["DID"])}），资金占用无显著变化（{cs(f11[2]["DID"])}）。')
pw, pw16, slope = DD['event_pre_wald'], DD['event_pre_wald_2016'], DD['event_pre_slope']
tr_ = DD['trend_adj']; fk = DD['placebo_fake']
def plev(p):
    return '在5%水平上显著' if p < 0.05 else ('仅在10%水平上显著' if p < 0.1 else '不显著')
body(f'列(4)、(5)以现金持有为被解释变量，PledgePre×Post的系数分别为{cs(m12[0]["coef"])}与{cs(m12[2]["coef"])}，均不显著；'
     f'剔除2018年（{m12[3]["coef"][0]}）、改用二值处理变量（{m12[4]["coef"][0]}）的结果相同。图3报告了事件研究的动态系数。'
     f'政策前（2007—2015年）系数的联合检验χ²({pw["df"]})={pw["statistic"]:.2f}（p={pw["pvalue"]:.3f}），{plev(pw["pvalue"])}；'
     f'纳入2016年后为χ²({pw16["df"]})={pw16["statistic"]:.2f}（p={pw16["pvalue"]:.3f}），{plev(pw16["pvalue"])}；'
     f'政策前系数的线性斜率为{slope["slope"]:.4f}（p={slope["p"]:.3f}），{plev(slope["p"])}；仅用政策前样本估计的线性趋势为{cs(DD["pre_only_trend"])}，{plev(DD["pre_only_trend_p"])}。'
     f'因此，平行趋势在5%水平上未被拒绝，但点估计在2007—2011年偏高、此后逐步下降，存在一定的事前趋势迹象。'
     f'加入处理强度与线性时间趋势的交乘项后，双重差分系数为{cs(tr_[0]["DID"])}（加入行业×年度与省份×年度固定效应后为{cs(tr_[1]["DID"])}），对设定较为敏感。'
     f'虚构政策时点为2014年、2015年的安慰剂系数分别为{cs(fk[0]["coef"])}与{cs(fk[1]["coef"])}；'
     f'在估计样本公司内随机置换处理强度{num(pp["n"])}次，绝对值不小于真实估计值的占{pct(pp["k_abs_ge_true"] / pp["n"])}（置换p值为{pp["p_two_sided"]:.3f}）。')
body('综合来看，新规实施后高质押公司的平仓压力相对上升，但质押比例的回落与均值回归幅度相当，资金占用与现金持有均未发生显著的相对变化。'
     '这不属于事先设定的任何一种情形，本文不将其解读为新规改变了质押风险向现金政策的传导，只作描述性报告。')
cap('表10　2018年股票质押新规的双重差分检验')
rows = [['变量', '(1)', '(2)', '(3)', '(4)', '(5)'], ['被解释变量', 'Pledge', 'Pressure', 'OREC', 'Cash', 'Cash'],
        ['PledgePre×Post'] + [c['DID'][0] for c in f11] + [m12[0]['coef'][0], m12[2]['coef'][0]],
        [''] + [c['DID'][1] for c in f11] + [m12[0]['coef'][1], m12[2]['coef'][1]],
        ['控制变量'] + yes(5), ['公司固定效应'] + yes(5), ['年度固定效应', '是', '是', '是', '是', '否'],
        ['行业×年度、省份×年度固定效应', '否', '否', '否', '否', '是'],
        ['调整R²'] + [c['r2'] for c in f11] + [m12[0]['r2'], m12[2]['r2']],
        ['观测值'] + [c['N'] for c in f11] + [m12[0]['N'], m12[2]['N']]]
table(rows, WD5, 2)
note('注：PledgePre为2016—2017年大股东平均质押比例，Post在2018年及以后取1；样本为2016年之前上市的公司。省级纾困计划的交错双重差分所需的各省纾困基金设立时间尚待依据毛捷和管星华（2022）原文补充。')
figp = add(clone(P_CAP, ''))
for r in figp.findall(W + 'r'):
    figp.remove(r)
Paragraph(figp, c4).add_run().add_picture(os.path.join(HERE, 'figures', 'fig5-1_event.png'), width=Cm(14.5))
add(clone(P_FIGCAP, '图3　2018年股票质押新规的动态效应'))
note('注：圆点为PledgePre×年度虚拟变量的系数，竖线为95%置信区间（公司层面聚类），2017年为基期；控制变量与固定效应同表10列(4)。')

# ---- 4. 稳健性与内生性 ----
sub('4．稳健性检验与内生性处理')
r16, r17, stl, alt = RB['t5_16'], RB['t5_17'], RB['stale'], RB['alt_construction']
lg, psm, psmu, iv, ivn = RB['t5_19_lag'], RB['t5_19_psm'], RB['t5_19_psm_unweighted'], RB['t5_19_iv'], RB['iv_n_cells']
w18, ss18, g18 = RB['t5_18_within'], RB['t5_18_samesample_firmfe'], RB['t5_18']
cx = RB['t5_15_capex']
rows = [['检验', 'Pledge系数', '标准误', '观测值'], ['基准（表6列(2)）', t6r[1]['coef'][0], t6r[1]['coef'][1], t6r[1]['N']]]
def rr(lab, c, n=None):
    rows.append([lab, c['coef'][0], c['coef'][1], n or c['N']])
rr('解释变量：是否存在质押', r16['Pledge_Dum']); rr('解释变量：占总股本质押比例', r16['Pledge_Ratio2']); rr('解释变量：明细表口径', r16['Pledge_det'])
rr('改进口径（非名义第一大股东、宽松名称匹配等）', alt)
rr('陈旧余额且十大股东无质押标识者置零', stl['zero_stale_flagno']); rr('陈旧余额全部置零', stl['zero_all_stale']); rr('剔除陈旧余额观测', stl['drop_stale'])
rr('剔除2008、2015、2020年', r17['drop_crisis']); rr('剔除上市前三年', r17['drop_first3'])
rr('公司与行业×年度双向聚类', r17['twoway']); rr('产权性质×年度固定效应', r17['soe_year']); rr('行业×年度＋产权性质×年度固定效应', r17['ind_soe_year'])
rr('公司×上市阶段固定效应（阶段内变化）', w18)
rr('质押比例滞后一期', lg); rr('PSM（对照组按匹配次数加权）', psm); rr('PSM（对照组去重、不加权）', psmu)
rr(f'2SLS（公司聚类，第一阶段F={iv["firm"]["F"]:.2f}）', iv['firm']['second']); rr(f'2SLS（行业聚类，第一阶段F={iv["industry"]["F"]:.2f}）', iv['industry']['second'])
rows.append(['被解释变量为资本支出Capex', cx['coef'][0], cx['coef'][1], cx['N']])
for key in ['Pledge_Dum', 'Pledge_Ratio2', 'Pledge_det']:
    need(not ns(r16[key]['coef']) and neg(r16[key]['coef']), '替换质押口径显著为负 ' + key)
for c in [r17['drop_crisis'], r17['drop_first3'], r17['twoway'], lg, psm, psmu, alt]:
    need(not ns(c['coef']) and neg(c['coef']), '稳健性显著为负')
need(not ns(stl['zero_all_stale']['coef']), '陈旧余额置零后仍显著')
need(ns(g18['4-6']['coef']) and ns(g18['7+']['coef']), '上市4—6年与7年以上组不显著')
body(f'表11汇总了稳健性检验与内生性处理的结果。替换质押口径、调整样本、改用双向聚类、质押比例滞后一期以及倾向得分匹配之后，Pledge的系数均保持显著为负。'
     f'有三点需要如实说明。第一，{pct(stl["share_among_pledgers"])}的质押观测仅由三年以前的统计表记录支撑，将其全部置零后系数为{stl["zero_all_stale"]["coef"][0]}，'
     f'绝对值有所下降但仍{sig(stl["zero_all_stale"]["coef"])}。第二，国有与非国有企业的质押比例时间趋势差异较大，允许二者有各自的年度效应后，'
     f'系数降为{cs(r17["soe_year"]["coef"])}，{sig(r17["soe_year"]["coef"])}；再加入行业×年度固定效应后为{cs(r17["ind_soe_year"]["coef"])}，{sig(r17["ind_soe_year"]["coef"])}。'
     f'这说明基准结果中有一部分可能来自不同产权性质企业在质押与现金持有上的不同时间趋势，正文将以此作为结论的重要限定。'
     f'第三，以同行业同年度（剔除本公司）平均质押比例为工具变量的2SLS系数为{iv["firm"]["second"]["coef"][0]}，约为OLS估计的'
     f'{round(float(iv["firm"]["second"]["coef"][0].rstrip("*")) / float(t6r[1]["coef"][0].rstrip("*")))}倍；工具变量仅在{num(ivn["industry_years"])}个行业—年度单元上变化，'
     f'按行业聚类时第一阶段F统计量为{iv["industry"]["F"]:.2f}，存在弱工具变量问题，加之排他性约束较弱，该结果仅作辅助参考。')
body(f'关于公司生命周期，按上市年限分组时，上市第4—6年与第7年及以后两组的系数分别为{g18["4-6"]["coef"][0]}与{g18["7+"]["coef"][0]}，均不显著，'
     f'上市第1—3年组为{g18["1-3"]["coef"][0]}，但该组有效观测仅{g18["1-3"]["N"]}个（现金流波动性要求连续三年数据，上市当年与次年的观测基本不在样本中）；'
     f'以公司×上市阶段固定效应只利用同一阶段内的变化时，系数为{cs(w18["coef"])}，约为同一样本公司固定效应估计（{ss18["coef"][0]}）的'
     f'{round(100 * float(w18["coef"][0].rstrip("*")) / float(ss18["coef"][0].rstrip("*")))}%。因此，负向关系有相当部分存在于同一上市阶段之内，'
     f'但上市初期募集资金的逐年消耗仍可能构成部分替代解释，须以IPO募集资金数据直接检验。拓展分析中，质押对资本支出的影响为{cs(cx["coef"])}，{sig(cx["coef"])}。')
cap('表11　稳健性检验与内生性处理汇总')
table(rows, [4300, 1600, 1500, 1600])
note('注：被解释变量为Cash（末行除外）；除注明者外，均控制表6列(2)的控制变量及公司与年度固定效应，括号内为公司层面聚类稳健标准误。'
     '“陈旧余额”指所有正余额出质方的最后一条统计表记录均早于年末3年以上；“上市阶段”分为上市第1—3年、第4—6年、第7年及以后；'
     'PSM为逐年logit估计倾向得分后1:1有放回近邻匹配（卡尺0.05）。')
# ---- 5. 小结 ----
need(stars(r17['soe_year']['coef']) < stars(t6r[1]['coef']), '产权×年度设定下显著性减弱')
sub('5．初步结论与后续研究的调整')
body('综合以上结果：第一，质押比例与现金持有之间存在稳健的负向关系，符号与H1b一致，但在允许不同产权性质企业有各自年度效应时显著性明显减弱；'
     '第二，平仓压力与资金占用两项渠道检验均未获支持，负向关系目前既不能归因于风险规避，也不能归因于以资金占用为表现的掏空，'
     '现金边际价值检验待数据补充后完成；第三，融资约束的调节作用实为上市年限的作用，产权性质的调节作用不显著，H3、H4均不成立；'
     '第四，2018年质押新规未带来现金持有的显著相对变化。后续研究将据此调整重点：补充IPO募集资金数据检验生命周期解释，'
     '补充现金边际价值检验以完成机制识别，并在正文中把产权性质×年度设定作为基准结论的重要稳健性检验。')

# ======================= 四(四)：尚需补充的数据 =======================
add(clone(P_SEC, '（四）尚需补充的数据'))
items = ['① 现金边际价值检验（模型4）：非经常性损益、利息费用、研发支出、普通股现金股利，以及筹资活动现金流中的吸收投资、取得借款、发行债券与偿还债务收付的现金；',
         '② KZ／WW 指数：现金股利、长期借款、应付债券；',
         '③ 被解释变量计量可靠性：使用受限的货币资金、利息收入、有息负债、货币资金相关违规处罚记录，“存贷双高”判定口径依杨国超等（2025）原文；',
         '④ 投资效率（Richardson，2006）与研发投入：研发支出、固定资产折旧与无形资产摊销、处置长期资产收回的现金、取得子公司支付的现金；',
         '⑤ 公司生命周期：IPO募集资金净额（用于检验上市后募集资金消耗对基准结果的影响）；',
         '⑥ 现金持有替代口径Cash2：交易性金融资产；',
         '⑦ 省级纾困计划：各省（区、市）地方政府纾困基金设立时间，依毛捷和管星华（2022）原文整理。']
for t in items:
    add(clone(P_ITEM, t))

# ======================= 五、进度安排 =======================
p = find_p(c5, lambda t: t.startswith('本文的主要工作进展及安排如表8所示'))
replace_in(p, '表8', '表12')
p = find_p(c5, lambda t: t.startswith('表8　'))
replace_in(p, '表8', '表12')
pt = [ch for ch in c5._tc.iterchildren() if ch.tag == W + 'tbl'][0]
NEWTASK = {'第一阶段': '数据下载、清洗与匹配；完成基准回归与U型关系复现（已完成）',
           '第二阶段': '变量测算（含SA、平仓压力）、描述性统计；渠道检验、边界条件、2018年新规DID、稳健性与内生性检验（已完成初步估计）',
           '第三阶段': '补充下载第四部分（四）所列数据，完成现金边际价值、KZ／WW、被解释变量可靠性、投资效率与生命周期检验',
           '第四阶段': '省级纾困补充检验；按导师与开题专家意见修订实证设计并重新估计'}
for tr in pt.findall(W + 'tr'):
    tcs = row_cells(tr)
    k = ptxt(tcs[0]).strip()
    if k in NEWTASK:
        set_tc(tcs[2], NEWTASK[k])

# 表题、图片段落与下一段同页（避免标题与表格分页）
def keep_next(pel):
    pPr = pel.find(W + 'pPr')
    if pPr is None:
        pPr = etree.SubElement(pel, W + 'pPr'); pel.remove(pPr); pel.insert(0, pPr)
    if pPr.find(W + 'keepNext') is None:
        kn = etree.Element(W + 'keepNext')
        pPr.insert(1 if pPr.find(W + 'pStyle') is not None else 0, kn)
for c in (c2, c4, c5):
    for pel in c._tc.iter(W + 'p'):
        t = ptxt(pel)
        if re.match(r'^表\d+　', t) or (pel.find('.//' + W + 'drawing') is not None and not t.strip()):
            keep_next(pel)
doc.save(OUT)
print('written', OUT)
